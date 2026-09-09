import io
import json
import pandas as pd
from pathlib import Path
from docx import Document
from email.message import EmailMessage

sample_dir = Path("sample_data")
sample_dir.mkdir(exist_ok=True)

# 1. Sample CSV
csv_data = """Name,Email,Date of Birth,Phone,Salary,City
RAVI KUMAR, ravi.kumar@example.com, 12/05/2001, +91-98765 43210, "₹10,000", Chennai
RAVI KUMAR, ravi.kumar@example.com, 12/05/2001, +91-98765 43210, "₹10,000", Chennai
EMILY BLUNT , emily@cinema.org , 1983-02-23 , +1 (212) 555-0199 , $125,000 , London
JOHN DOE, NA, 04/18/1992, 9876500000, $85000, New York
"""
with open(sample_dir / "sample_users.csv", "w", encoding="utf-8") as f:
    f.write(csv_data)

# 2. Sample Excel
df = pd.DataFrame([
    {"Customer Name": "ALEXANDER HAMILTON", "Transaction Date": "04/07/2026", "Email Address": "alex@treasury.gov", "Amount Paid": "$4,500.50", "Status": "Completed"},
    {"Customer Name": "THOMAS JEFFERSON", "Transaction Date": "13/04/1743", "Email Address": "thomas@monticello.org", "Amount Paid": "$8,200.00", "Status": "Completed"},
    {"Customer Name": "BENJAMIN FRANKLIN", "Transaction Date": "17/01/1706", "Email Address": "ben@poorrichard.com", "Amount Paid": "₹15,000", "Status": "Pending"},
    {"Customer Name": "BENJAMIN FRANKLIN", "Transaction Date": "17/01/1706", "Email Address": "ben@poorrichard.com", "Amount Paid": "₹15,000", "Status": "Pending"},
])
df.to_excel(sample_dir / "sample_transactions.xlsx", index=False)

# 3. Sample JSON
json_data = [
    {"name": "VICTORIA SECRET", "email": "VICTORIA@EXAMPLE.COM", "phone": "+44 20 7946 0912", "dob": "1994-11-30", "total": "£550.00"},
    {"name": "VICTORIA SECRET", "email": "VICTORIA@EXAMPLE.COM", "phone": "+44 20 7946 0912", "dob": "1994-11-30", "total": "£550.00"},
    {"name": "MICHAEL CORLEONE", "email": "michael@corleone.it", "phone": "+1 (212) 555-7890", "dob": "1950-03-24", "total": "$10,000"}
]
with open(sample_dir / "sample_records.json", "w", encoding="utf-8") as f:
    json.dump(json_data, f, indent=2)

# 4. Sample Unstructured TXT
txt_data = """CUSTOMER ONBOARDING FORM

Name: Ravi Kumar
Date of Birth: 12/05/2001
Email: ravi@gmail.com
Phone: +91-98765 43210
Address: 42 Marina Bay Road, Chennai
Postal Code: 600001
ID Number: IND-8823-9011
Total Amount: INR 10,000
Organization: DataFlow Technologies Inc.
"""
with open(sample_dir / "sample_profile.txt", "w", encoding="utf-8") as f:
    f.write(txt_data)

# 5. Sample DOCX
doc = Document()
doc.add_heading('Employee Verification Document', level=1)
doc.add_paragraph('Employee Name: Sarah Jenkins')
doc.add_paragraph('Date of Birth: 1988-06-15')
doc.add_paragraph('Email: sarah.jenkins@acme-corp.com')
doc.add_paragraph('Phone: +1 (415) 555-4321')
doc.add_paragraph('City: San Francisco')
doc.add_paragraph('Postal Code: 94107')
doc.add_paragraph('Total Amount: $12,500.00')
doc.save(sample_dir / "sample_verification.docx")

# 6. Sample EML
msg = EmailMessage()
msg["From"] = "Arthur Dent <arthur.dent@galaxy.org>"
msg["To"] = "Support <support@dataflow.com>"
msg["Subject"] = "Account Update Request"
msg["Date"] = "Wed, 09 Sep 2026 09:30:00 +0000"
msg.set_content("""
Hello DataFlow Team,

Please update my contact information with the details below:

Name: Arthur Dent
Date of Birth: 1978-03-11
Phone: +44 1632 960999
City: London
Postal Code: SW1A 1AA
Total Amount: £42.00

Kind regards,
Arthur Dent
""")
with open(sample_dir / "sample_email.eml", "wb") as f:
    f.write(msg.as_bytes())

print("Sample files successfully generated in 'sample_data/' folder!")
