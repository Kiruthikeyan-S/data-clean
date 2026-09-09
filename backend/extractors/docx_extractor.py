import io
from docx import Document
from typing import Dict, Any, List

def extract_docx(file_bytes: bytes) -> Dict[str, Any]:
    """
    Extracts text from Word .docx documents including paragraphs and tables.
    """
    doc_io = io.BytesIO(file_bytes)
    doc = Document(doc_io)
    
    extracted_parts: List[str] = []
    
    # 1. Paragraphs
    for p in doc.paragraphs:
        text = p.text.strip()
        if text:
            extracted_parts.append(text)
            
    # 2. Tables
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
            if row_text:
                extracted_parts.append(row_text)
                
    full_text = "\n".join(extracted_parts)
    return {
        "text": full_text,
        "confidence": 1.0,
        "extraction_method": "docx_parser"
    }
