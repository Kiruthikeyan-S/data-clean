import sys
import io
import json
import pandas as pd
from PIL import Image, ImageDraw
import pymupdf as fitz
from docx import Document
from email.message import EmailMessage
from backend.pipeline import process_file_pipeline

if sys.platform == "win32":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

def test_all():
    print("=== 1. Testing CSV ===")
    csv_content = """Name,Email,Date of Birth,Phone,Salary
ALICE SMITH , ALICE@EXAMPLE.COM, 1990-01-15, +1-555-0199, $85000
ALICE SMITH , ALICE@EXAMPLE.COM, 1990-01-15, +1-555-0199, $85000
Bob Jones, bob@test, NA, 9876543210, "120,000"
"""
    res_csv = process_file_pipeline("employees.csv", "text/csv", csv_content.encode("utf-8"))
    assert res_csv.classification == "structured"
    assert len(res_csv.structured_data) == 2
    print(f"CSV OK: {len(res_csv.structured_data)} records")

    print("\n=== 2. Testing Excel ===")
    df = pd.DataFrame([
        {"Product": "Laptop", "Price": "₹75,000", "Stock": 10},
        {"Product": "Mouse", "Price": "$25", "Stock": 100},
        {"Product": "Mouse", "Price": "$25", "Stock": 100}
    ])
    excel_buf = io.BytesIO()
    df.to_excel(excel_buf, index=False)
    res_excel = process_file_pipeline("inventory.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", excel_buf.getvalue())
    assert res_excel.classification == "structured"
    assert len(res_excel.structured_data) == 2
    print(f"Excel OK: {len(res_excel.structured_data)} records")

    print("\n=== 3. Testing JSON ===")
    json_bytes = json.dumps([
        {"name": " Charlie Brown ", "email": "CHARLIE@TEST.COM", "phone": "9998887776"},
        {"name": " Charlie Brown ", "email": "CHARLIE@TEST.COM", "phone": "9998887776"}
    ]).encode("utf-8")
    res_json = process_file_pipeline("users.json", "application/json", json_bytes)
    assert res_json.classification == "structured"
    assert len(res_json.structured_data) == 1
    print(f"JSON OK: {len(res_json.structured_data)} records")

    print("\n=== 4. Testing TXT ===")
    txt_content = """
    Customer Statement
    Name: Sarah Connor
    DOB: 1984-02-28
    Email: sarah.connor@sky.net
    Phone: +1 415 555 2671
    City: Los Angeles
    Total Amount: $4,500.00
    """
    res_txt = process_file_pipeline("statement.txt", "text/plain", txt_content.encode("utf-8"))
    assert res_txt.classification == "unstructured"
    fields_map = {f.key: f.value for f in res_txt.fields}
    assert fields_map.get("name") == "Sarah Connor"
    assert fields_map.get("email") == "sarah.connor@sky.net"
    assert "total_amount" in fields_map or "amount" in fields_map
    print(f"TXT OK: {fields_map}")

    print("\n=== 5. Testing DOCX ===")
    doc = Document()
    doc.add_heading('Invoice Summary', level=1)
    doc.add_paragraph('Name: Bruce Wayne')
    doc.add_paragraph('Date of Birth: 1980-05-27')
    doc.add_paragraph('Email: bruce@wayne-enterprises.com')
    doc.add_paragraph('Phone: +1 212 555 0198')
    doc.add_paragraph('Total Amount: $1,000,000')
    docx_buf = io.BytesIO()
    doc.save(docx_buf)
    res_docx = process_file_pipeline("invoice.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", docx_buf.getvalue())
    assert res_docx.classification == "unstructured"
    fields_map = {f.key: f.value for f in res_docx.fields}
    assert fields_map["name"] == "Bruce Wayne"
    assert fields_map["email"] == "bruce@wayne-enterprises.com"
    print(f"DOCX OK: {fields_map}")

    print("\n=== 6. Testing EML ===")
    msg = EmailMessage()
    msg["From"] = "John Doe <john.doe@enterprise.com>"
    msg["To"] = "Support <support@dataflow.com>"
    msg["Subject"] = "Account Registration Request"
    msg["Date"] = "Wed, 09 Sep 2026 10:00:00 +0000"
    msg.set_content("""
    Dear Support Team,
    
    Please register the following user:
    Name: John Doe
    Date of Birth: 1995-10-20
    Phone: +91 9123456780
    City: Bangalore
    Postal Code: 560001
    
    Thanks!
    """)
    res_eml = process_file_pipeline("registration.eml", "message/rfc822", msg.as_bytes())
    assert res_eml.classification == "unstructured"
    fields_map = {f.key: f.value for f in res_eml.fields}
    assert any(k in fields_map for k in ["name", "user_name", "full_name", "sender"])
    print(f"EML OK: {fields_map}")

    print("\n=== 7. Testing PDF Native ===")
    doc_pdf = fitz.open()
    page = doc_pdf.new_page()
    page.insert_text((50, 72), "Medical Report\nName: Peter Parker\nDOB: 2001-08-10\nEmail: peter.parker@dailybugle.com\nPhone: +1 212 555 9876\nTotal Amount: $150.00")
    pdf_bytes = doc_pdf.tobytes()
    doc_pdf.close()
    res_pdf = process_file_pipeline("report.pdf", "application/pdf", pdf_bytes)
    assert res_pdf.classification == "unstructured"
    fields_map = {f.key: f.value for f in res_pdf.fields}
    assert fields_map["name"] == "Peter Parker"
    assert fields_map["email"] == "peter.parker@dailybugle.com"
    print(f"PDF Native OK: {fields_map}")

    print("\n=== 8. Testing Image OCR ===")
    img = Image.new("RGB", (600, 250), color=(255, 255, 255))
    d = ImageDraw.Draw(img)
    d.text((30, 30), "Name: Clark Kent", fill=(0, 0, 0))
    d.text((30, 70), "Email: clark@dailyplanet.com", fill=(0, 0, 0))
    d.text((30, 110), "Phone: +1 312 555 0144", fill=(0, 0, 0))
    d.text((30, 150), "Total: $250.00", fill=(0, 0, 0))
    img_buf = io.BytesIO()
    img.save(img_buf, format="PNG")
    res_img = process_file_pipeline("card.png", "image/png", img_buf.getvalue())
    assert res_img.classification == "unstructured"
    fields_map = {f.key: f.value for f in res_img.fields}
    print(f"Image OCR OK: {fields_map}")

    print("\n==========================================")
    print("ALL 8 FORMATS TESTED AND PASSED EXCELLENTLY!")
    print("==========================================")

if __name__ == "__main__":
    test_all()
